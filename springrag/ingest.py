"""Ingesta: convierte archivos (HTML, PDF, DOCX, TXT) en texto limpio + metadatos.

Cada loader devuelve una lista de `Document`. Un Document es un trozo de
texto con los metadatos necesarios para citar su origen más adelante.
HTML se parte por secciones (h1/h2/h3) y PDF por páginas, para que el
chunking posterior no mezcle temas distintos.
"""
from dataclasses import dataclass, field
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxFile
from pypdf import PdfReader

from springrag.sources import Source


@dataclass
class Document:
    text: str
    metadata: dict = field(default_factory=dict)


def _meta(src: Source | None, path: Path, **extra) -> dict:
    """Metadatos comunes. Si no viene del catálogo se usa el nombre del archivo."""
    return {
        "source_title": src.title if src else path.stem,
        "url": src.url if src else str(path),
        "doc_type": path.suffix.lstrip("."),
        **extra,
    }


def _inside(el, names: tuple) -> bool:
    """¿El elemento está dentro de alguna de estas etiquetas? (evita texto duplicado)"""
    return any(p.name in names for p in el.parents)


def load_html(path: Path, src: Source | None = None) -> list[Document]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")

    # Nos quedamos con el contenido principal y quitamos menús, scripts, etc.
    main = soup.find("article") or soup.find("main") or soup.body
    for tag in main.find_all(["script", "style", "nav", "header", "footer", "aside"]):
        tag.decompose()

    docs: list[Document] = []
    section, parts = "Introducción", []

    def cerrar_seccion():
        texto = "\n".join(parts).strip()
        if texto:
            docs.append(Document(texto, _meta(src, path, section=section)))

    for el in main.find_all(["h1", "h2", "h3", "p", "pre", "li", "table"]):
        if el.name in ("h1", "h2", "h3"):
            cerrar_seccion()  # un encabezado nuevo = sección nueva
            section, parts = el.get_text(" ", strip=True), []
        elif el.name == "table":
            filas = [" | ".join(c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"]))
                     for tr in el.find_all("tr")]
            parts.append("\n".join(filas))
        elif _inside(el, ("table",)) or (el.name != "li" and _inside(el, ("li",))):
            continue  # ya quedó incluido dentro de su tabla o de su <li>
        elif el.name == "li" and _inside(el, ("li",)):
            continue  # <li> anidado: lo incluye su <li> padre
        elif el.name == "pre":
            parts.append("```\n" + el.get_text().strip() + "\n```")
        else:
            parts.append(el.get_text(" ", strip=True))
    cerrar_seccion()
    return docs


def load_pdf(path: Path, src: Source | None = None) -> list[Document]:
    reader = PdfReader(str(path))
    docs = []
    for n, page in enumerate(reader.pages, start=1):
        texto = (page.extract_text() or "").strip()
        if texto:
            docs.append(Document(texto, _meta(src, path, section=f"Página {n}", page=n)))
    return docs


def load_docx(path: Path, src: Source | None = None) -> list[Document]:
    texto = "\n".join(p.text for p in DocxFile(str(path)).paragraphs if p.text.strip())
    return [Document(texto, _meta(src, path, section="Documento"))]


def load_txt(path: Path, src: Source | None = None) -> list[Document]:
    texto = path.read_text(encoding="utf-8", errors="ignore").strip()
    return [Document(texto, _meta(src, path, section="Documento"))]


LOADERS = {".html": load_html, ".pdf": load_pdf, ".docx": load_docx, ".txt": load_txt}


def load_file(path: Path, src: Source | None = None) -> list[Document]:
    """Elige el loader según la extensión del archivo."""
    return LOADERS[path.suffix.lower()](path, src)
